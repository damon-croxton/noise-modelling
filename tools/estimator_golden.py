#!/usr/bin/env python3
"""Generate golden test cases from the TfNSW Construction Noise Estimator.

    python3 tools/estimator_golden.py <TfNSW_Calculator.xlsx> [tests/golden/estimator-cases.json]

Needs openpyxl and LibreOffice (soffice) on the developer's machine; users never
run this. For each case the spreadsheet is filled in, recalculated by
LibreOffice and the outputs read back:
  Step 2  C14 cumulative SWL (and the rows it ranks)
  Step 3  Q7/I7/J7/K7 levels, F16-J16 NMLs, K17 sleep screening, F19-J19
          exceedances
  Step 4  the additional measures flagged for every period and band.

LibreOffice 24.2 has no XLOOKUP, used only to read the quantity of the
loudest rows (Step 2, G18:G20). Those three cells are replaced by the
equivalent INDEX/MATCH before recalculation; nothing else is changed.

The spreadsheet itself is TfNSW material and is not stored in the repository.
"""
import json, os, shutil, subprocess, sys, tempfile

import openpyxl

CASES = []
# A. Plant combinations (Step 2 rule) with assorted propagation and NML inputs.
#    plant: [name in the estimator list, qty, intermittent, locally screened]
A = [
    ("one item", [("Excavator (20 tonne)", 1, False, False)], 40, 0, 0.0, "Oct to April", "R3", {}),
    ("loudest x3, 5 dB above next -> 1 row", [("Excavator (20 tonne)", 3, False, False), ("Light vehicles", 1, False, False)], 60, 0, 0.5, "May to Sep", "R2", {}),
    ("loudest x2 -> 2 rows", [("Jackhammer*", 2, False, False), ("Generator - Diesel", 1, False, False), ("Hand Power Tool", 1, False, False)], 25, 5, 0.2, "Oct to April", "R4", {}),
    ("loudest x3 within 5 dB -> 2 rows", [("Concrete Pump", 3, False, False), ("Concrete Truck / Agitator", 1, False, False), ("Elevated Work Platform", 1, False, False)], 100, 0, 1.0, "Oct to April", "R1", {"night": 28}),
    ("loudest x1, second x2 -> 2 rows", [("Concrete saw/corer*", 1, False, False), ("Excavator (10 tonne)", 2, False, False), ("Light vehicles (eg 4WD)", 1, False, False)], 50, 10, 0.5, "May to Sep", "R5", {}),
    ("three single rows + quieter fourth -> 3 rows", [("Piling Rig - Bored", 1, False, False), ("Mobile Crane (20 tonne)", 1, False, False), ("Generator - Diesel", 1, False, False), ("Welder", 1, False, False)], 150, 15, 0.8, "Oct to April", "R3", {"day": 47, "evening": 41, "night": 36}),
    ("intermittent and screened flags", [("Jackhammer*", 1, True, False), ("Concrete saw/corer*", 1, True, True), ("Excavator (20 tonne)", 1, False, True)], 30, 0, 0.0, "May to Sep", "R4", {"day": 33, "evening": 29, "night": 25}),
    ("tie between rows", [("Excavator (20 tonne)", 1, False, False), ("Auger/Drill Rig", 1, False, False), ("Welder", 1, False, False)], 80, 5, 0.3, "Oct to April", "R2", {}),
    ("reference-case plant", [("Scissor Lift", 1, True, False), ("Impact wrench", 1, True, False), ("Hand Power Tool", 1, True, False), ("Light vehicles", 1, True, False)], 35, 0, 0.0, "Oct to April", "R4", {"night": 39}),
    ("large crew, single type x6", [("Hand Power Tools (2-3 Items)", 6, False, False), ("Light vehicles", 2, False, False)], 20, 0, 0.0, "May to Sep", "R5", {}),
]
for name, plant, d, scr, soft, season, area, rbl in A:
    CASES.append({"kind": "plant", "name": name, "plant": plant, "dist": d, "screening": scr, "soft": soft, "season": season, "area": area, "rbl": rbl})
# B. Exceedances on and around the band boundaries: total SWL set directly so
#    the standard-hours exceedance at 10 m (R3, NML 60) is x; the other periods
#    sit 5 / 10 / ~15 dB higher because their NMLs are lower.
for x in [-1, -0.6, 0, 0.4, 0.5, 4, 5, 9, 10, 14, 15, 19, 20]:
    CASES.append({"kind": "exceedance", "name": f"standard-hours exceedance {x}", "swl": 60 + x + 28 + 2.98666666666667 * 10 / 1000,
                  "dist": 10, "screening": 0, "soft": 0.0, "season": "Oct to April", "area": "R3", "rbl": {}})
CASES.append({"kind": "exceedance", "name": "75 dBA highly noise affected", "swl": 75.2 + 28 + 0.0299, "dist": 10, "screening": 0, "soft": 0.0, "season": "Oct to April", "area": "R5", "rbl": {}})

STEP4_ROWS = {13: "PN", 14: "V", 15: "SN", 16: "RO", 17: "RP", 18: "DR", 19: "AA"}
STEP4_COLS = {"std": "CDEF", "p1day": "GHIJ", "p1eve": "KLMN", "p2eve": "OPQR", "p2night": "STUV"}


def fill(src, case, out):
    wb = openpyxl.load_workbook(src)
    s2, s3 = wb["Step 2 Cumulative_Level"], wb["Step 3 Results"]
    names = {" ".join(str(s2[f"B{r}"].value).split()): r for r in range(22, 164) if s2[f"B{r}"].value}
    rows = []
    if case["kind"] == "plant":
        for nm, qty, inter, scr in case["plant"]:
            r = names[" ".join(nm.split())]
            s2[f"D{r}"] = qty; s2[f"E{r}"] = "Yes" if inter else "No"; s2[f"F{r}"] = "Yes" if scr else "No"
            rows.append({"name": nm, "row": r, "lw": s2[f"C{r}"].value, "qty": qty, "intermittent": inter, "screened": scr})
        rng = "$G$22:$G$163", "$D$22:$D$163"
        s2["G18"] = f"=INDEX({rng[1]},MATCH(MAX(G22:G163),{rng[0]},0))"
        s2["G19"] = f"=INDEX({rng[1]},MATCH(LARGE(G22:G163,2),{rng[0]},0))"
        s2["G20"] = f"=INDEX({rng[1]},MATCH(LARGE(G22:G163,3),{rng[0]},0))"
    else:
        s3["E7"] = case["swl"]
    s3["D7"] = case["season"]; s3["F7"] = case["dist"]; s3["G7"] = case["screening"]; s3["H7"] = case["soft"]
    s3["C12"] = case["area"]
    for cell, k in (("F14", "day"), ("H14", "day"), ("I14", "evening"), ("J14", "night")):
        s3[cell] = case["rbl"].get(k)
    wb.save(out)
    return rows


def main():
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), "..", "tests", "golden", "estimator-cases.json")
    tmp = tempfile.mkdtemp(prefix="estgold-")
    ind, outd, prof = (os.path.join(tmp, d) for d in ("in", "out", "profile"))
    os.makedirs(ind); os.makedirs(outd); os.makedirs(os.path.join(prof, "user"))
    with open(os.path.join(prof, "user", "registrymodifications.xcu"), "w") as f:   # recalculate on load
        f.write('<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item></oor:items>')
    plant_rows = []
    for i, c in enumerate(CASES):
        plant_rows.append(fill(src, c, os.path.join(ind, f"c{i:02d}.xlsx")))
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outd]
                   + sorted(os.path.join(ind, f) for f in os.listdir(ind)), check=True, capture_output=True)
    out = []
    for i, c in enumerate(CASES):
        wb = openpyxl.load_workbook(os.path.join(outd, f"c{i:02d}.xlsx"), data_only=True)
        s2, s3, s4 = wb["Step 2 Cumulative_Level"], wb["Step 3 Results"], wb["Step 4 Mitigation Measures"]
        v = lambda ws, a: ws[a].value
        res = {
            "swl": v(s3, "E7"), "Q7": v(s3, "Q7"), "I7": v(s3, "I7"), "J7": v(s3, "J7"), "K7": v(s3, "K7"),
            "nml": {"std": v(s3, "F16"), "p1day": v(s3, "H16"), "p1eve": v(s3, "I16"), "p2eve": v(s3, "I16"), "p2night": v(s3, "J16")},
            "sleepScreen": v(s3, "K17"), "sleepExceeded": v(s3, "K19"),
            "exceedance": {"std": v(s3, "F19"), "p1day": v(s3, "H19"), "p1eve": v(s3, "I19"), "p2eve": v(s3, "I19"), "p2night": v(s3, "J19")},
            "measures": {p: sorted({code for r, code in STEP4_ROWS.items() for col in cols if v(s4, f"{col}{r}") == "Yes"}) for p, cols in STEP4_COLS.items()},
        }
        if c["kind"] == "plant":
            res["ranked"] = [v(s2, "E18"), v(s2, "E19"), v(s2, "E20")]
            res["rankedQty"] = [v(s2, "G18"), v(s2, "G19"), v(s2, "G20")]
        out.append(dict(c, plantRows=plant_rows[i], result=res))
    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    with open(dst, "w") as f:
        json.dump({"_about": "Generated by tools/estimator_golden.py from the TfNSW Construction Noise Estimator (EMF-NV-TT-0068), recalculated in LibreOffice. Do not edit by hand.",
                   "cases": out}, f, indent=1)
    shutil.rmtree(tmp)
    print(f"wrote {len(out)} cases to {dst}")


if __name__ == "__main__":
    main()
